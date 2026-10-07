"""Lê uma apresentação PPTX existente e gera um RASCUNHO de deck.json (conteúdo preservado, estilo noir).
Uso: python extract_pptx.py <arquivo.pptx> <pasta_de_trabalho>
Saída: <pasta>/deck.json + <pasta>/extracted/ (imagens). Depois REVISE o deck.json slide a slide (títulos como
frases-conclusão, tipo certo de slide, no máx. 3 painéis e 5 KPIs) e rode noir_pipeline.py.
Heurísticas: 1º slide = cover; slide só com título = divider; último com 'obrigado/thanks/perguntas' = closing;
gráficos viram painéis line/bars/donut; tabelas viram painel table; imagens viram painel image; texto vira bullets/text."""
import sys, os, re, json
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE_TYPE

CLOSING = re.compile(r"obrigad|thank|perguntas|questions|contato|contact", re.I)


def shape_texts(sh):
    return [p.text.strip() for p in sh.text_frame.paragraphs if p.text.strip()] if sh.has_text_frame else []


def biggest_size(sh):
    sizes = [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size]
    return max(sizes) if sizes else 0


def chart_panel(ch, title):
    try:
        pl = ch.plots[0]; cats = [str(c) for c in pl.categories]; ser = pl.series[0]; vals = [float(v or 0) for v in ser.values]
    except Exception:
        return None
    t = str(ch.chart_type).lower()
    if "pie" in t or "doughnut" in t: return {"type": "donut", "title": title, "parts": [[c, v] for c, v in zip(cats, vals)]}
    if "line" in t or "area" in t: return {"type": "line", "title": title, "cats": cats, "vals": vals}
    return {"type": "bars", "title": title, "cats": cats, "vals": vals}


def extract(path, work):
    prs = Presentation(path); out = os.path.join(work, "extracted"); os.makedirs(out, exist_ok=True)
    slides = []; n = len(prs.slides._sldIdLst)
    for i, s in enumerate(prs.slides, 1):
        title = None; texts = []; panels = []; img = 0
        shapes = sorted(s.shapes, key=lambda x: (x.top or 0, x.left or 0))
        if s.shapes.title is not None and s.shapes.title.has_text_frame and s.shapes.title.text.strip(): title = s.shapes.title.text.strip()
        if not title:
            cand = [(biggest_size(sh), sh) for sh in shapes if sh.has_text_frame and shape_texts(sh)]
            if cand: title = shape_texts(max(cand, key=lambda c: c[0])[1])[0]
        for sh in shapes:
            if sh.has_text_frame and not (s.shapes.title is not None and sh.shape_id == s.shapes.title.shape_id):
                t = shape_texts(sh)
                if t and t != [title]: texts.append(t)
            if getattr(sh, "has_chart", False) and sh.has_chart:
                p = chart_panel(sh.chart, (sh.chart.chart_title.text_frame.text if sh.chart.has_title else "Gráfico"))
                if p: panels.append(p)
            if getattr(sh, "has_table", False) and sh.has_table:
                rows = [[c.text.strip() for c in r.cells] for r in sh.table.rows]
                if rows: panels.append({"type": "table", "title": "Tabela", "headers": rows[0], "rows": rows[1:20]})
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                img += 1; fn = os.path.join(out, f"slide{i}_img{img}.{sh.image.ext}"); open(fn, "wb").write(sh.image.blob)
                panels.append({"type": "image", "title": "Imagem", "src": fn})
        flat = [l for t in texts for l in t]
        title = title or f"Slide {i}"
        if i == 1:
            slides.append({"type": "cover", "title": title[:28], "overlay": (flat[0] if flat else "")[:40], "subtitle": " · ".join(flat[1:3]), "object": "auto",
                           "meta": [[str(k + 1), t] for k, t in enumerate(flat[3:7])], "_original_title": title})
        elif i == n and CLOSING.search(" ".join([title] + flat)):
            slides.append({"type": "closing", "title": title[:24].upper(), "overlay": "", "lines": flat[:3], "object": "auto"})
        elif not panels and len(flat) <= 1 and len(title) <= 40:
            slides.append({"type": "divider", "title": title, "overlay": flat[0] if flat else "", "object": "auto"})
        else:
            if flat:
                if len(flat) == 1 and len(flat[0]) > 120: panels.append({"type": "text", "title": "Resumo", "body": flat[0]})
                else: panels.append({"type": "bullets", "title": "Pontos-chave", "items": flat[:8]})
            for k in range(0, max(len(panels), 1), 3):
                chunk = panels[k:k + 3]
                slides.append({"type": "content", "kicker": "", "title": title + (" (cont.)" if k else ""), "panels": chunk})
    deck = {"meta": {"title": slides[0].get("_original_title", "Apresentação"), "org": "", "lang": "pt-BR", "accent": "#F7941D", "logo": None,
                     "topbar": "", "footer": "", "objects": [], "_nota": "REVISAR: org, logo, acento, objetos 3D (assets/objects ou gerar no Canva) e todos os títulos"}, "slides": slides}
    json.dump(deck, open(os.path.join(work, "deck.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(len(slides), "slides no rascunho ->", os.path.join(work, "deck.json"))


if __name__ == "__main__":
    os.makedirs(sys.argv[2], exist_ok=True); extract(sys.argv[1], sys.argv[2])
