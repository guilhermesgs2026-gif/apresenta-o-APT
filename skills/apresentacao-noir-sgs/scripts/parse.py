import json
from pptx import Presentation
p=Presentation("fonte.pptx")
out=[]
for i,s in enumerate(p.slides,1):
    d={"n":i,"texts":[],"charts":[],"tables":[]}
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip(): d["texts"].append(sh.text_frame.text.strip())
        if getattr(sh,"has_chart",False) and sh.has_chart:
            c=sh.chart; pl=c.plots[0]
            d["charts"].append({"type":str(c.chart_type),"cats":list(pl.categories),"vals":list(pl.series[0].values)})
        if getattr(sh,"has_table",False) and sh.has_table:
            d["tables"].append([[c.text for c in r.cells] for r in sh.table.rows])
    out.append(d)
json.dump(out,open("data_raw.json","w",encoding="utf-8"),ensure_ascii=False,indent=1)
for d in out[:7]+out[31:]: print(d["n"],d["texts"][:14])
